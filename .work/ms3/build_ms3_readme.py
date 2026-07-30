from __future__ import annotations

from copy import deepcopy
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SOURCE = Path("/Users/chesterlim/Desktop/Orbital 26/MS2/TheCSPlan-MS2-README.docx")
WORK_DIR = Path("/Users/chesterlim/Desktop/Orbital 26/The CS Plan/.work/ms3")
OUTPUT = Path("/Users/chesterlim/Desktop/Orbital 26/The CS Plan/MS3 README.docx")
ERD_PATH = WORK_DIR / "ms3-erd.png"
PLACEHOLDER_DIR = WORK_DIR / "placeholders"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    name = "Arial Bold.ttf" if bold else "Arial.ttf"
    return ImageFont.truetype(f"/System/Library/Fonts/Supplemental/{name}", size)


def create_placeholder(path: Path, title: str, subtitle: str, split: bool = False) -> None:
    width, height = 1800, 620
    image = Image.new("RGB", (width, height), "#f7f7f7")
    draw = ImageDraw.Draw(image)
    border = "#8a8a8a"
    for offset in range(0, width - 80, 32):
        draw.line((40 + offset, 40, min(56 + offset, width - 40), 40), fill=border, width=3)
        draw.line((40 + offset, height - 40, min(56 + offset, width - 40), height - 40), fill=border, width=3)
    for offset in range(0, height - 80, 32):
        draw.line((40, 40 + offset, 40, min(56 + offset, height - 40)), fill=border, width=3)
        draw.line((width - 40, 40 + offset, width - 40, min(56 + offset, height - 40)), fill=border, width=3)
    if split:
        for offset in range(80, height - 80, 32):
            draw.line((width // 2, offset, width // 2, min(offset + 16, height - 80)), fill="#b0b0b0", width=3)
        draw.text((width // 4, 152), "LEFT SCREENSHOT", font=font(38, True), fill="#616161", anchor="mm")
        draw.text((3 * width // 4, 152), "RIGHT SCREENSHOT", font=font(38, True), fill="#616161", anchor="mm")
    draw.text((width // 2, height // 2 - 36), title, font=font(48, True), fill="#333333", anchor="mm")
    draw.text((width // 2, height // 2 + 42), subtitle, font=font(29), fill="#666666", anchor="mm")
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def create_erd(path: Path) -> None:
    width, height = 3000, 1740
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    body_fill = "#f6f7f8"
    header_fill = "#e7eaee"
    stroke = "#222222"
    title_font = font(32, True)
    body_font = font(24)
    label_font = font(23)

    boxes = {
        "User": (80, 70, 600, 400),
        "StudentProfile": (820, 70, 1500, 510),
        "PlanExport (API payload)": (1760, 70, 2920, 350),
        "AdminCurriculumDraft (API payload)": (1760, 430, 2920, 790),
        "Module": (80, 560, 600, 1080),
        "Plan": (820, 630, 1500, 1060),
        "SemesterPlan": (1800, 890, 2320, 1180),
        "PlanItem": (1800, 1280, 2920, 1630),
        "ModuleRequirementTags": (80, 1310, 600, 1680),
        "RequirementSet": (820, 1310, 1500, 1680),
    }
    fields = {
        "User": [
            "_id",
            "googleId",
            "email",
            "name",
            "avatarUrl?",
            "role: student | admin",
            "createdAt, updatedAt",
        ],
        "StudentProfile": [
            "_id",
            "userId -> User",
            "programme",
            "cohort",
            "startingSemester",
            "currentSemester",
            "graduationSemester",
            "primaryPlanId? -> Plan",
            "planOrder[] -> Plan",
            "createdAt, updatedAt",
        ],
        "PlanExport (API payload)": [
            "schemaVersion: 1",
            "exportedAt",
            "profile: StudentProfile snapshot",
            "plan: Plan snapshot without MongoDB ids",
            "encoded into tcp1_ payload or import link",
        ],
        "AdminCurriculumDraft (API payload)": [
            "programme, cohort",
            "baseVersion -> latest RequirementSet",
            "totalUnits",
            "sourceNote",
            "rules[]: RequirementRule",
            "tagChanges[] -> ModuleRequirementTags",
            "local until validated and published",
        ],
        "Module": [
            "_id",
            "acadYear",
            "moduleCode",
            "title",
            "units",
            "department?",
            "faculty?",
            "description?",
            "prerequisite?",
            "prereqTree?",
            "createdAt, updatedAt",
        ],
        "Plan": [
            "_id",
            "userId -> User",
            "name",
            "programme",
            "cohort",
            "semesters[]",
            "createdAt, updatedAt",
        ],
        "SemesterPlan": [
            "key",
            "label",
            "items[]",
        ],
        "PlanItem": [
            "ModulePlanItem:",
            "id?, type: module, moduleCode -> Module, units,",
            "status, grade?, isSu",
            "",
            "PlaceholderPlanItem:",
            "id?, type: placeholder, requirementId, label,",
            "units, grade?, isSu",
        ],
        "ModuleRequirementTags": [
            "_id",
            "programme",
            "cohort",
            "moduleCode -> Module",
            "tags[]",
            "createdAt, updatedAt",
        ],
        "RequirementSet": [
            "_id",
            "programme",
            "cohort",
            "version",
            "totalUnits",
            "sourceNote",
            "rules[]: RequirementRule",
            "createdAt, updatedAt",
        ],
    }

    def draw_box(name: str) -> None:
        x1, y1, x2, y2 = boxes[name]
        header_height = 52
        draw.rectangle((x1, y1, x2, y2), fill=body_fill, outline=stroke, width=3)
        draw.rectangle((x1, y1, x2, y1 + header_height), fill=header_fill, outline=stroke, width=3)
        draw.text((x1 + 18, y1 + 10), name, font=title_font, fill=stroke)
        current_y = y1 + header_height + 22
        for line in fields[name]:
            draw.text((x1 + 24, current_y), line, font=body_font, fill=stroke)
            current_y += 31

    def connector(points: list[tuple[int, int]], label: str, label_xy: tuple[int, int]) -> None:
        draw.line(points, fill=stroke, width=3, joint="curve")
        for endpoint in (points[0], points[-1]):
            x, y = endpoint
            draw.ellipse((x - 7, y - 7, x + 7, y + 7), fill=stroke)
        draw.text(label_xy, label, font=label_font, fill=stroke)

    for name in boxes:
        draw_box(name)

    connector([(600, 205), (710, 205), (710, 190), (820, 190)], "1 to 1 profile", (640, 154))
    connector([(600, 350), (635, 350), (635, 820), (820, 820)], "1 to many plans", (650, 545))
    connector([(1160, 510), (1160, 630)], "primaryPlanId / planOrder", (1190, 555))
    connector([(1500, 820), (1560, 820), (1560, 1010), (1800, 1010)], "semesters[] contains", (1580, 865))
    connector([(2060, 1180), (2060, 1280)], "items[] contains", (2100, 1215))
    connector([(600, 870), (690, 870), (690, 1225), (1690, 1225), (1690, 1450), (1800, 1450)], "moduleCode reference", (735, 1185))
    connector([(340, 1080), (340, 1310)], "tagged by", (190, 1190))
    connector([(600, 1490), (820, 1490)], "rules use tags", (635, 1452))
    connector([(1160, 1060), (1160, 1310)], "evaluated against", (1195, 1168))
    connector([(1500, 200), (1760, 200)], "snapshot", (1578, 158))

    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)


def find_paragraph(doc: Document, text: str):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def replace_paragraph_text(paragraph, text: str) -> None:
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    if paragraph.style.name == "Heading 1":
        run.bold = True
        run.font.size = Pt(18)
    elif paragraph.style.name == "Heading 2":
        run.underline = True


def delete_between(start_paragraph, end_paragraph) -> None:
    current = start_paragraph._p.getnext()
    while current is not None and current is not end_paragraph._p:
        next_element = current.getnext()
        current.getparent().remove(current)
        current = next_element


def delete_after_to_section_end(paragraph) -> None:
    current = paragraph._p.getnext()
    while current is not None and current.tag != qn("w:sectPr"):
        next_element = current.getnext()
        current.getparent().remove(current)
        current = next_element


def set_paragraph_format(paragraph, role: str) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_with_next = role in {"h1", "h2", "h3"}
    if role == "h1":
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(0)
    elif role in {"h2", "h3"}:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(10)
    else:
        paragraph.paragraph_format.space_after = Pt(6)


def add_paragraph(doc: Document, anchor, text: str, role: str = "body", italic: bool = False):
    style = {
        "h1": "Heading 1",
        "h2": "Heading 2",
        "h3": "Heading 3",
        "body": "normal",
        "caption": "normal",
    }[role]
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    if role == "h1":
        run.bold = True
        run.font.size = Pt(18)
    elif role == "h2":
        run.underline = True
    if italic:
        run.italic = True
    set_paragraph_format(paragraph, role)
    if role == "caption":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(10)
        run.italic = True
    anchor._p.addprevious(paragraph._p)
    return paragraph


def set_cell_border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")


def set_cell_margins(cell, top: int = 70, start: int = 100, bottom: int = 70, end: int = 100) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        element = tc_mar.find(qn(f"w:{tag}"))
        if element is None:
            element = OxmlElement(f"w:{tag}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def add_table(doc: Document, anchor, headers: tuple[str, str], rows: Iterable[tuple[str, str]]):
    row_values = list(rows)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.columns[0].width = Inches(2.85)
    table.columns[1].width = Inches(3.35)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "8895")
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    def fill_row(cells, values: tuple[str, str], header: bool = False) -> None:
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(2.85 if index == 0 else 3.35)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles["normal"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.bold = header

    def set_row_properties(row, repeat_header: bool = False) -> None:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = tr_pr.find(qn("w:cantSplit"))
        if cant_split is None:
            cant_split = OxmlElement("w:cantSplit")
            tr_pr.append(cant_split)
        cant_split.set(qn("w:val"), "1")
        if repeat_header:
            tbl_header = tr_pr.find(qn("w:tblHeader"))
            if tbl_header is None:
                tbl_header = OxmlElement("w:tblHeader")
                tr_pr.append(tbl_header)
            tbl_header.set(qn("w:val"), "1")

    fill_row(table.rows[0].cells, headers, header=True)
    set_row_properties(table.rows[0], repeat_header=True)
    for row in row_values:
        added_row = table.add_row()
        fill_row(added_row.cells, row)
        set_row_properties(added_row)

    anchor._p.addprevious(table._tbl)
    spacer = doc.add_paragraph(style="normal")
    spacer.paragraph_format.space_after = Pt(0)
    anchor._p.addprevious(spacer._p)
    return table


def add_picture(doc: Document, anchor, path: Path, width: float = 6.4):
    paragraph = doc.add_paragraph(style="normal")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    anchor._p.addprevious(paragraph._p)
    return paragraph


def add_feature_table(doc: Document, anchor, rows: list[tuple[str, str]]) -> None:
    add_table(doc, anchor, ("Important file", "Purpose"), rows)


def build() -> None:
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    create_erd(ERD_PATH)
    placeholders = {
        "view-rules": ("Insert Admin Dashboard VIEW screenshot", "Selected curriculum with the full ruleset and expandable rule JSON", False),
        "view-tags": ("Insert module mapping VIEW screenshots", "Left: module search and tag filter // Right: complete module-tag inventory", True),
        "edit-rules": ("Insert Admin Dashboard EDIT screenshots", "Left: rule JSON editor and unit fields // Right: release diff, validation, and publish controls", True),
        "edit-tags": ("Insert staged mapping EDIT screenshots", "Left: staged tag update // Right: deleted mapping and local diff", True),
        "create-clone": ("Insert CREATE and CLONE screenshots", "Left: new curriculum draft form // Right: clone preview and confirmation", True),
        "catalog-validation": ("Insert published-curriculum validation screenshots", "Left: onboarding programme/cohort options // Right: settings programme/cohort options", True),
        "five-programmes": ("Insert five-programme scaling screenshots", "Programme selector and representative degree requirement pages for the supported curricula", True),
    }
    placeholder_paths: dict[str, Path] = {}
    for key, (title, subtitle, split) in placeholders.items():
        path = PLACEHOLDER_DIR / f"{key}.png"
        create_placeholder(path, title, subtitle, split)
        placeholder_paths[key] = path

    doc = Document(SOURCE)

    milestone = find_paragraph(doc, "Milestone 2")
    replace_paragraph_text(milestone, "Milestone 3")
    milestone.runs[0].bold = True
    milestone.runs[0].font.size = Pt(26)

    last_existing_future_story = find_paragraph(
        doc,
        "As a maintainer, I want requirement changes to be validated against sample plans before publishing so that scaling to more degrees does not break existing CS requirement evaluation.",
    )
    future_story_anchor = Paragraph(
        last_existing_future_story._p.getnext(),
        last_existing_future_story._parent,
    )
    add_paragraph(doc, future_story_anchor, "Future Plans", "h3")
    add_paragraph(
        doc,
        future_story_anchor,
        "As an NUS student outside the School of Computing, I want to choose my home faculty, programme, and cohort so that the planner evaluates my academic plan against the correct faculty-owned ruleset.",
    )
    add_paragraph(
        doc,
        future_story_anchor,
        "As a faculty curriculum maintainer, I want curriculum drafts to record their source, reviewer, approval history, and rollback target so that published requirement changes are traceable and recoverable.",
    )
    add_paragraph(
        doc,
        future_story_anchor,
        "As a student with unfilled semesters, I want the planner to recommend modules that satisfy my remaining requirements while respecting prerequisites, module availability, timetable clashes, and workload limits so that I can turn degree progress into a feasible semester plan.",
    )

    data_flow_anchor = find_paragraph(doc, "Data and Application Flow")

    add_paragraph(doc, data_flow_anchor, "8. Admin Dashboard: View Curriculum Rules and Module Mappings", "h2")
    add_paragraph(doc, data_flow_anchor, "Description", "h3")
    add_paragraph(
        doc,
        data_flow_anchor,
        "The administrator dashboard is a role-protected workspace for inspecting every published curriculum. An administrator first selects a programme and cohort from the curriculum catalogue. Each catalogue entry shows the latest version and opens the complete RequirementSet for that curriculum, including its total unit target, source note, ordered requirement rules, and current version number. This makes the database representation visible in the web application instead of requiring direct MongoDB inspection or reading seeded TypeScript data.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "Every requirement is displayed as a rule card and can be expanded to show the rule's complete JSON. This is important because the rule type alone is not enough to understand a curriculum: the administrator must also see required units, accepted tags, module choices, tag caps, placeholder categories, focus areas, industry constraints, honours pathways, and any advisory text. Viewing the canonical JSON keeps the dashboard faithful to the same RequirementRule shape that is evaluated by the rules engine.",
    )
    add_picture(doc, data_flow_anchor, placeholder_paths["view-rules"])
    add_paragraph(
        doc,
        data_flow_anchor,
        "Admin dashboard curriculum view showing the selected programme, cohort, latest version, total units, source note, ordered rules, and expandable rule JSON.",
        "caption",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "The Module Tags tab exposes the entire programme-and-cohort-specific module mapping. Without a query, it pages through all mapped modules; with a query, it searches the module catalogue by module code or title and joins the matching modules to their current tags. Administrators can also filter the results by one tag. A separate tag inventory lists every tag used by the curriculum together with the modules mapped to it, allowing the administrator to inspect the classification vocabulary before editing a rule or mapping.",
    )
    add_picture(doc, data_flow_anchor, placeholder_paths["view-tags"])
    add_paragraph(
        doc,
        data_flow_anchor,
        "Left: searchable module mappings with a module-tag filter // Right: the complete tag inventory for the selected curriculum.",
        "caption",
    )
    add_feature_table(
        doc,
        data_flow_anchor,
        [
            ("apps/web/src/pages/AdministratorPage.tsx", "Curriculum selector, requirement cards, JSON view, mapping search, pagination, tag filtering, and tag inventory."),
            ("apps/api/src/routes/adminRoutes.ts", "Admin-only read routes for the curriculum catalogue, latest RequirementSet, module mappings, and available tags."),
            ("apps/api/src/services/adminCurriculumService.ts", "Aggregates version counts and mapping counts, loads the latest ruleset, searches modules, and groups the complete tag inventory."),
            ("apps/api/src/middleware/auth.ts", "requireAdmin rejects non-administrator accounts before any administration route runs."),
            ("apps/api/src/models/RequirementSet.ts", "Stores versioned curriculum rules selected by programme and cohort."),
            ("apps/api/src/models/ModuleRequirementTags.ts", "Stores programme/cohort/module-specific classification tags."),
            ("apps/api/src/models/Module.ts", "Provides the searchable NUS module catalogue joined to curriculum mappings."),
        ],
    )
    add_paragraph(doc, data_flow_anchor, "Implementation Philosophy", "h3")
    add_paragraph(
        doc,
        data_flow_anchor,
        "The view workflow is deliberately read-only and authoritative. Curriculum rules and module mappings are loaded through dedicated admin endpoints guarded by requireAuth and requireAdmin, while the frontend only stores navigation state such as the selected curriculum, active tab, search, filter, and page. This separation prevents a viewing action from becoming an accidental database mutation and makes the latest published version the clear source of truth.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "Module metadata and curriculum meaning remain separate. Module stores general NUS information, while ModuleRequirementTags classifies the same module differently for each programme and cohort. The dashboard preserves that model by searching the shared module catalogue first and then presenting the mapping in the selected curriculum context. It therefore scales without copying module title, description, prerequisite, or faculty data into every ruleset.",
    )

    add_paragraph(doc, data_flow_anchor, "9. Admin Dashboard: Edit, Validate, Diff, and Publish", "h2")
    add_paragraph(doc, data_flow_anchor, "Description", "h3")
    add_paragraph(
        doc,
        data_flow_anchor,
        "The edit workflow lets an administrator change a curriculum's total units, source note, rule order, and individual requirement rules directly in the web application. Each rule can be edited as formatted JSON. The editor parses the text and validates it with RequirementRuleSchema before replacing the in-memory rule, so temporarily incomplete JSON remains local to the editor and invalid rule objects never enter the curriculum draft. Required units and every rule-specific threshold remain part of that same canonical JSON representation.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "New rules are added through a rule id, label, and rule-type selector. Selecting a type creates the matching JSON template instead of starting from an unstructured empty object. For example, a module-list rule starts with requiredModules, a capped-units-from-tags rule starts with acceptedTags and tagCaps, and the structured rule types start with the threshold arrays and unit fields needed by their evaluator. The administrator can then edit, reorder, duplicate, or remove rules while preserving the explicit priority order used during allocation.",
    )
    add_picture(doc, data_flow_anchor, placeholder_paths["edit-rules"])
    add_paragraph(
        doc,
        data_flow_anchor,
        "Left: editing one requirement's JSON and unit fields in the web application // Right: rule and mapping diff, validation state, next version, and publish controls.",
        "caption",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "Module mappings use a staged editing model. An administrator searches for a module, edits its comma-separated tag keys, and stages the result as an AdminModuleTagChange. A mapping can also be staged for deletion. Re-entering the published tags removes the staged change, and the complete staged set can be cleared without affecting the database. Rule edits and mapping edits are persisted in administrator-specific localStorage so that changing tabs or returning to the dashboard does not lose an unfinished draft.",
    )
    add_picture(doc, data_flow_anchor, placeholder_paths["edit-tags"])
    add_paragraph(
        doc,
        data_flow_anchor,
        "Left: a locally staged module-tag update // Right: a staged mapping deletion with the corresponding diff before publication.",
        "caption",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "The release sidebar continuously produces a diff against the latest published RequirementSet. It reports total-unit and source-note changes, rules added, changed, or removed, ordinary module-tag changes, mapping deletions, and the next version number. Validation then checks the entire draft rather than a single field: shared Zod schemas, unique rule ids, required fields for each rule type, known module references, duplicate tag changes, referenced tags, declared unit totals, the latest base version, and regression previews against up to five saved plans for the same programme and cohort.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "All edits remain local until the exact current draft has passed validation and the administrator confirms publication. Publishing creates a new RequirementSet version and applies all staged mapping upserts or deletions in one MongoDB transaction. The previous RequirementSet remains stored, and a baseVersion check rejects stale drafts if another version was published after the draft was loaded. After success, React Query invalidates the administration, requirements, module-tag, and evaluation caches so users receive the new curriculum data.",
    )
    add_feature_table(
        doc,
        data_flow_anchor,
        [
            ("apps/web/src/pages/AdministratorPage.tsx", "Rule JSON editor, type-tagged rule templates, local draft persistence, staged mapping changes, diff summary, validation fingerprint, and publish controls."),
            ("packages/shared/src/index.ts", "RequirementRuleSchema, RequirementRuleTypeSchema, AdminModuleTagChangeSchema, and AdminCurriculumDraftSchema."),
            ("apps/api/src/services/adminCurriculumService.ts", "Semantic validation, tag coverage checks, sample-plan regression previews, base-version checks, transactional publication, and mapping bulk writes."),
            ("apps/api/src/routes/adminRoutes.ts", "POST /validate and POST /publish endpoints behind administrator authorization."),
            ("apps/api/src/models/RequirementSet.ts", "Preserves every published version under the unique programme/cohort/version identity."),
            ("apps/api/src/models/ModuleRequirementTags.ts", "Receives staged mapping upserts and deletions only during publication."),
            ("packages/rules-engine/src/evaluatePlan.ts", "Evaluates saved sample plans against the proposed RequirementSet during validation."),
        ],
    )
    add_paragraph(doc, data_flow_anchor, "Implementation Philosophy", "h3")
    add_paragraph(
        doc,
        data_flow_anchor,
        "Curriculum maintenance follows a draft-validate-publish workflow rather than in-place editing. Local drafts make experimentation safe, the diff makes the release understandable, and versioned publication keeps prior rulesets available for diagnosis. The validated fingerprint prevents the administrator from validating one draft, changing it, and publishing a different draft without another validation pass.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "Rules and module mappings are published atomically because they describe one curriculum interpretation. Creating a RequirementSet version without its matching tag changes could make valid rules evaluate incorrectly, while changing tags without the intended rules could silently reallocate modules. A MongoDB transaction and optimistic baseVersion check prevent those split-brain and lost-update states.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "Rule-type templates are intentionally explicit. Each evaluator branch expects a different combination of fields, so selecting a rule type creates a valid structural starting point and exposes the resulting JSON for review. This keeps the admin interface flexible enough for advanced curricula without hiding the data model behind dozens of one-off form screens.",
    )

    add_paragraph(doc, data_flow_anchor, "10. Admin Dashboard: Create and Clone Curricula", "h2")
    add_paragraph(doc, data_flow_anchor, "Description", "h3")
    add_paragraph(
        doc,
        data_flow_anchor,
        "An administrator can create an entirely new curriculum by selecting a programme identity and entering a cohort. CohortSchema requires the AY2025/26 format and verifies that the end year is exactly the year after the start year. The dashboard also checks the published catalogue so an existing programme/cohort pair cannot be recreated accidentally. A new curriculum begins as a local baseVersion 0 draft with a 160-unit default, an empty source note, one starter rule, and no module-tag changes; it becomes visible to students only after full validation and publication as version 1.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "The clone workflow is optimised for a new academic year whose requirements have not changed. The administrator selects a complete source curriculum and a target programme/cohort, then previews the operation. The preview shows the source version, total units, rule count, tag count, and source note; checks whether the target already has any RequirementSet version or module mappings; preserves references to historical modules while warning about modules missing from the current catalogue; and reports tags referenced by rules but absent from the source mappings.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "If the target is clear, cloning copies the entire rule array, source note, total units, and all module mappings into a transaction and creates the target at version 1. This is a complete curriculum copy rather than a frontend convenience: both the RequirementSet and ModuleRequirementTags records receive the new programme/cohort identity, so the cloned cohort is immediately independent and can be edited and versioned later.",
    )
    add_picture(doc, data_flow_anchor, placeholder_paths["create-clone"])
    add_paragraph(
        doc,
        data_flow_anchor,
        "Left: creating a new local curriculum draft with validated programme and cohort // Right: clone preview showing source statistics, conflicts, warnings, and the target version.",
        "caption",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "Onboarding and Settings use the published requirement-set catalogue as their allowed programme and cohort options. The programme dropdown is derived from programmes that have at least one published ruleset, and the cohort dropdown is filtered to published cohorts for the selected programme. A user therefore cannot choose an unsupported curriculum through these pages, while a newly published or cloned curriculum becomes selectable without hardcoding another frontend option.",
    )
    add_picture(doc, data_flow_anchor, placeholder_paths["catalog-validation"])
    add_paragraph(
        doc,
        data_flow_anchor,
        "Left: onboarding only offers published programme/cohort combinations // Right: Settings applies the same curriculum-backed validation before changing all plans.",
        "caption",
    )
    add_feature_table(
        doc,
        data_flow_anchor,
        [
            ("apps/web/src/pages/AdministratorPage.tsx", "New-curriculum form, cohort validation feedback, clone target selection, preview, confirmation, and local baseVersion 0 drafts."),
            ("packages/shared/src/index.ts", "CohortSchema validates consecutive academic years; AdminCloneCurriculumSchema prevents cloning onto the same curriculum."),
            ("apps/api/src/services/adminCurriculumService.ts", "Clone preview, conflict detection, warning generation, and transactional copying of rules and module mappings."),
            ("apps/api/src/services/requirementService.ts", "Lists published programme/cohort pairs used as the student-facing curriculum catalogue."),
            ("apps/web/src/pages/OnboardingPage.tsx", "Builds programme and cohort choices only from published rulesets."),
            ("apps/web/src/pages/SettingsPage.tsx", "Restricts academic profile changes to published programme/cohort combinations."),
            ("apps/api/src/services/profileService.ts", "Applies the selected published curriculum identity to the profile and every owned plan."),
        ],
    )
    add_paragraph(doc, data_flow_anchor, "Implementation Philosophy", "h3")
    add_paragraph(
        doc,
        data_flow_anchor,
        "Creation and cloning share the same publication boundary. A new identifier is not treated as supported merely because it appears in the programme enum; it becomes a usable curriculum only when a valid RequirementSet exists. This makes the database catalogue a capability registry for onboarding, settings, evaluation, and administration.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "Cloning is intentionally exact and conservative. If requirements are unchanged, copying the complete versioned curriculum is safer than re-entering rules and mappings. If requirements later diverge, the administrator edits the independent target draft and publishes a new version. Target conflict checks and transactions prevent a clone from overwriting an existing cohort or leaving a half-created curriculum.",
    )

    add_paragraph(doc, data_flow_anchor, "11. Scaling to Five School of Computing Curricula", "h2")
    add_paragraph(doc, data_flow_anchor, "Description", "h3")
    add_paragraph(
        doc,
        data_flow_anchor,
        "Milestone 3 scales The CS Plan from one Computer Science ruleset to five live programme identities: Computer Science, Business Analytics, Business Artificial Intelligence Systems, Computer Science and Mathematics Double Major, and Business Analytics and Economics Double Degree. Computer Science is available for AY2025/26 and AY2026/27, so the five programmes are represented by six published programme/cohort rulesets. Each curriculum has its own total-unit target, official source note, ordered requirement rules, and module-tag mappings.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "The rulesets were created by translating the published curriculum structures into RequirementSet data and mapping relevant module codes into programme/cohort-specific tags. Existing rule types were reused wherever they represented the published requirement accurately: module-list, units-from-tags, capped-units-from-tags, structured-idcd, structured-breadth-depth, and residual-units continue to serve multiple programmes. New evaluator branches were added only when a genuinely new structure could not be expressed safely with existing fields, including module-choice, structured-programme-electives, structured-industry-experience, and structured-ddp-honours-pathway.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "Business Analytics and Business Artificial Intelligence Systems reuse the same general University Pillars, Computing Ethics, ID/CD, programme-elective, industry-experience, and unrestricted-elective concepts while supplying different core modules, tags, thresholds, and dissertation or internship alternatives. Their mapping generators also classify broad families of NUS modules, then assign unmapped catalogue modules to unrestricted electives where the curriculum permits.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "Computer Science and Mathematics Double Major reuses the AY2025/26 Computer Science ruleset and module mappings as its base. It replaces the Mathematics and Sciences rule with the second-major categories and caps, removes conflicting UE or CD classifications from Mathematics modules, and reduces the unrestricted-elective target. This demonstrates that a related curriculum can reuse stable rules and alter only the parts that differ.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "A double major or double degree is modelled as one curriculum identity, not as two independently selected programmes. The Business Analytics and Economics Double Degree therefore has one dedicated requirement set that encodes shared courses, both discipline cores, programme electives, and the integrated-thesis or internship honours pathways. Treating the published course as one curriculum avoids ambiguous merging, duplicate allocation, and mismatched total-unit targets.",
    )
    add_picture(doc, data_flow_anchor, placeholder_paths["five-programmes"])
    add_paragraph(
        doc,
        data_flow_anchor,
        "Five supported programme identities: Computer Science, Business Analytics, Business Artificial Intelligence Systems, Computer Science and Mathematics Double Major, and Business Analytics and Economics Double Degree.",
        "caption",
    )
    add_feature_table(
        doc,
        data_flow_anchor,
        [
            ("packages/data/src/computerScienceData.ts", "Computer Science AY2025/26 and AY2026/27 rulesets, focus areas, mappings, and generated common classifications."),
            ("packages/data/src/businessAnalyticsData.ts", "Business Analytics core, statistics alternatives, programme electives, industry pathways, and mappings."),
            ("packages/data/src/baisData.ts", "Business Artificial Intelligence Systems core, programme electives, industry pathways, and mappings."),
            ("packages/data/src/csMathDoubleMajData.ts", "Reuses Computer Science rules and mappings, then replaces Mathematics requirements and adjusts unrestricted electives."),
            ("packages/data/src/bzaEconsDdpData.ts", "Dedicated Business Analytics and Economics Double Degree curriculum and honours-pathway mappings."),
            ("packages/rules-engine/src/evaluatePlan.ts", "Shared evaluator branches reused across all five programme identities."),
            ("apps/api/src/upsertCurricula.ts", "Upserts the current curricula and complete generated module mappings into MongoDB."),
            ("apps/api/src/tests/*Data.test.ts", "Curriculum-specific data regression tests for totals, rules, official constraints, and mappings."),
        ],
    )
    add_paragraph(doc, data_flow_anchor, "Implementation Philosophy", "h3")
    add_paragraph(
        doc,
        data_flow_anchor,
        "Scaling is data-first. React pages and Express routes do not branch on every degree; they request the selected programme/cohort RequirementSet and mapping lookup, then pass both to the same evaluator. The user interface can therefore support a newly published curriculum without a new planner page, while genuinely new academic logic is introduced as a named, validated rule type rather than programme-specific conditionals.",
    )
    add_paragraph(
        doc,
        data_flow_anchor,
        "Curriculum identities match the choices published to students. A combined course can reuse rules and mappings internally, but it remains one selectable programme with one allocation model and one total-unit target. This prevents the planner from naively combining two standalone programmes whose overlapping modules, common curriculum, honours routes, or unrestricted-elective rules may interact.",
    )

    software_anchor = find_paragraph(doc, "Software Engineering Practices")
    add_paragraph(doc, software_anchor, "Milestone 3 Curriculum Administration Flow", "h2")
    add_paragraph(
        doc,
        software_anchor,
        "An authenticated administrator opens /administrator. The frontend verifies the User.role returned by /api/me, redirects non-administrators to the planner, and enables all /api/admin calls only for an administrator.",
    )
    add_paragraph(
        doc,
        software_anchor,
        "The dashboard loads the curriculum catalogue, the latest RequirementSet for the selected programme/cohort, the current module mappings, and the complete tag inventory. Curriculum, editor, and search state are namespaced to the administrator and curriculum in localStorage.",
    )
    add_paragraph(
        doc,
        software_anchor,
        "Rule JSON, total units, source note, order, and staged AdminModuleTagChange objects are edited locally. The frontend compares this draft with the latest published ruleset and displays rule additions, changes, removals, mapping changes, mapping deletions, and the next version.",
    )
    add_paragraph(
        doc,
        software_anchor,
        "Validation parses the AdminCurriculumDraft through shared Zod schemas, checks rule-specific semantics and module references, applies staged mappings to an in-memory lookup, warns about missing tags or mismatched unit totals, checks baseVersion, and evaluates up to five recent saved plans for that curriculum.",
    )
    add_paragraph(
        doc,
        software_anchor,
        "Publication revalidates the draft, starts a MongoDB session, confirms that baseVersion is still current, inserts version baseVersion + 1, and bulk-writes the staged mapping updates or deletions in the same transaction. The frontend then clears the local draft and invalidates curriculum-dependent caches.",
    )
    add_paragraph(
        doc,
        software_anchor,
        "Cloning follows a separate preview path. The backend verifies that the target has no RequirementSet or mappings, reports retained historical modules and unmapped tags, and then transactionally copies the complete source curriculum into target version 1.",
    )
    add_paragraph(
        doc,
        software_anchor,
        "Student onboarding and settings request the published curriculum catalogue. The selected programme filters the allowed cohorts, so only an existing RequirementSet can be written to StudentProfile and propagated to the user's plans.",
    )

    add_paragraph(doc, software_anchor, "Technology Stack", "h1")
    add_paragraph(doc, software_anchor, "Frontend", "h2")
    add_paragraph(
        doc,
        software_anchor,
        "The web application is written in TypeScript 5.8 with React 19.1 and React DOM. Vite 6.3 provides the development server and production bundle, while React Router 7.6 defines login, onboarding, planner, GPA, requirements, settings, and administrator routes. TanStack React Query 5.81 manages server state, caching, mutations, and invalidation across profile, plan, evaluation, curriculum, and administration workflows.",
    )
    add_paragraph(
        doc,
        software_anchor,
        "Forms use React Hook Form 7.58 with @hookform/resolvers and shared Zod 3.25 schemas. Tailwind CSS 3.4, PostCSS, Autoprefixer, clsx, and tailwind-merge implement the responsive dark interface and reusable class composition. Lucide React supplies the interface icon set. Low-risk client persistence uses localStorage for evaluation caches, dismissed warning state, and administrator drafts or editor state; authoritative academic data remains on the API.",
    )
    add_paragraph(doc, software_anchor, "Backend", "h2")
    add_paragraph(
        doc,
        software_anchor,
        "The API is a Node.js TypeScript application built with Express 5.1. Express routers separate authentication, current-user, profile, plan, module, requirement, and administration endpoints. Zod validates route bodies and shared domain objects. Mongoose 8.16 provides models, indexes, queries, aggregation pipelines, bulk writes, sessions, and transactions. The API uses JSON Web Tokens, cookie-parser, and HTTP-only tcp_session cookies for seven-day sessions; CORS is configured for credentialed requests from the web client.",
    )
    add_paragraph(
        doc,
        software_anchor,
        "Google OAuth is implemented in the backend without coupling Google tokens to academic records. dotenv loads deployment configuration such as MONGODB_URI, CLIENT_URL, SERVER_URL, SESSION_SECRET, GOOGLE_CLIENT_ID, and GOOGLE_CLIENT_SECRET. The backend normalises NUSMods catalogue data, serves owner-scoped plans, composes requirement evaluation inputs, and enforces the administrator role before curriculum operations.",
    )
    add_paragraph(doc, software_anchor, "Database and Domain Packages", "h2")
    add_paragraph(
        doc,
        software_anchor,
        "MongoDB stores User, StudentProfile, Plan, Module, ModuleRequirementTags, and RequirementSet collections. Plans embed semester and plan-item snapshots because planner and GPA updates operate on the complete scenario. RequirementSet documents are append-only by version, while programme/cohort/module mappings are updated transactionally with a publication. Compound unique indexes protect module catalogue identities, mapping identities, and curriculum versions.",
    )
    add_paragraph(
        doc,
        software_anchor,
        "The npm-workspace monorepo separates packages/shared for Zod schemas, types, programme labels, GPA constants, semester helpers, and S/U eligibility; packages/data for the five supported curricula and generated mappings; and packages/rules-engine for pure requirement allocation and prerequisite evaluation. This keeps React, Express, data definition, and curriculum evaluation independently testable while sharing one TypeScript contract.",
    )
    add_paragraph(doc, software_anchor, "Tooling, Testing, and Deployment", "h2")
    add_paragraph(
        doc,
        software_anchor,
        "npm workspaces coordinate builds and tests, tsx runs TypeScript development and data scripts, and concurrently starts the web and API applications together. Vitest 3.2 runs API, shared-domain, curriculum-data, and rules-engine tests; Supertest is available for HTTP-level API workflow coverage. TypeScript project configurations provide package builds and no-emit type checking. The frontend includes a Vercel single-page-application rewrite, while the API runs from apps/api with environment-based database and OAuth configuration.",
    )

    data_driven_paragraph = find_paragraph(
        doc,
        "Degree Requirement Fulfillment is implemented through packages/rules-engine instead of route-level or component-level conditionals. Requirements are represented as versioned data, while the evaluator handles allocation, placeholders, tag matching, structured ID/CD, breadth/depth rules, residual units, and warnings. This practice is essential for Milestone 3 because new programmes and cohorts should mostly require new requirement data rather than new UI branches.",
    )
    replace_paragraph_text(
        data_driven_paragraph,
        "Degree Requirement Fulfillment is implemented through packages/rules-engine instead of route-level or component-level conditionals. Requirements are represented as versioned data, while the evaluator handles allocation, placeholders, tag matching, structured ID/CD, breadth/depth rules, programme electives, industry experience, double-degree honours pathways, residual units, and warnings. This practice enabled Milestone 3 to add five programme identities and another cohort primarily through new requirement data and deliberate reusable rule types rather than new UI branches.",
    )

    testing_anchor = find_paragraph(doc, "Testing")
    add_paragraph(doc, testing_anchor, "Role-Based Access Control", "h2")
    add_paragraph(
        doc,
        testing_anchor,
        "The administrator feature applies the same authorization rule at both interface and API boundaries. AdministratorPage redirects a student account, while every /api/admin route passes through requireAuth and requireAdmin. This practice protects curriculum viewing, validation, publication, and cloning even if a non-administrator calls the API directly.",
    )
    add_paragraph(doc, testing_anchor, "Draft, Diff, Validation, and Publish Workflow", "h2")
    add_paragraph(
        doc,
        testing_anchor,
        "Curriculum editing is split into explicit states: published data, local draft, validated draft, and published next version. The feature uses administrator-and-curriculum-scoped localStorage, a human-readable diff, and a JSON fingerprint of the validated draft. Publication is disabled whenever the draft changes after validation. This practice links directly to rule editing, staged module mappings, deletions, and new-curriculum creation.",
    )
    add_paragraph(doc, testing_anchor, "Transactional Versioning and Optimistic Concurrency", "h2")
    add_paragraph(
        doc,
        testing_anchor,
        "RequirementSet publication and ModuleRequirementTags changes run in one MongoDB transaction. A baseVersion comparison is performed during validation and again inside the transaction, so concurrent administrators cannot silently overwrite one another. Every successful release creates a new RequirementSet version, while cloning creates a conflict-checked version 1 for a new curriculum identity.",
    )
    add_paragraph(doc, testing_anchor, "Rule-Type Templates and Deliberate Reuse", "h2")
    add_paragraph(
        doc,
        testing_anchor,
        "The administrator's rule-type selector is tagged to createRequirementRule, which generates the fields required by the chosen evaluator branch. Curriculum data reuses stable rule types and adds a new type only when a published constraint cannot be represented safely. This practice supports the JSON editor, five-programme scaling, and consistent semantic validation.",
    )
    add_paragraph(doc, testing_anchor, "Curriculum-Catalogue-Driven Interfaces", "h2")
    add_paragraph(
        doc,
        testing_anchor,
        "Onboarding and Settings derive programme and cohort choices from published RequirementSet identities. The same catalogue is used by the administrator selector and student-facing profile flows. This removes duplicated option lists and ensures that users can only choose a curriculum that the evaluator can load.",
    )
    add_paragraph(doc, testing_anchor, "Programme-Specific Module Classification", "h2")
    add_paragraph(
        doc,
        testing_anchor,
        "ModuleRequirementTags keeps shared NUS module metadata separate from curriculum-specific meaning. Generated and explicit mappings can therefore classify one module differently across Computer Science, Business Analytics, BAIS, a double major, and a double degree. This practice supports requirement evaluation, module-tag administration, and future expansion without duplicating the module catalogue.",
    )
    add_paragraph(doc, testing_anchor, "Pure Domain Logic and Regression Fixtures", "h2")
    add_paragraph(
        doc,
        testing_anchor,
        "Requirement evaluation, prerequisite checking, S/U eligibility, schema parsing, and admin semantic checks are deterministic functions with focused Vitest coverage. Curriculum-specific fixture tests assert published unit totals, rule structures, mappings, and complex pathways. This practice makes ruleset reuse and administration changes safer because regressions are caught before data is deployed.",
    )

    future_heading = find_paragraph(doc, "Upcoming Milestone 3 Features")
    delete_between(testing_anchor, future_heading)
    add_paragraph(
        doc,
        future_heading,
        "The Milestone 3 repository test suite contains 83 passing Vitest tests across 10 test files. It covers the rules engine, S/U eligibility, NUSMods normalisation, authentication and administrator guards, administrator draft contracts and semantic validation, and regression checks for all five supported programme identities. The suite is run across npm workspaces with npm test.",
    )
    add_paragraph(doc, future_heading, "Unit and Domain Logic Testing", "h2")
    add_paragraph(
        doc,
        future_heading,
        "packages/rules-engine/src/tests/evaluatePlan.test.ts contains 48 tests for allocation, module lists, module choices, tag caps, placeholders, ID/CD constraints, breadth and depth, programme electives, industry experience, double-degree honours routes, unrestricted-elective overflow, programme-specific mappings, unknown modules, and prerequisite-tree evaluation. apps/api/src/tests/suEligibility.test.ts adds seven GPA-page policy tests for module levels, prerequisite-bearing Level-2000 modules, published exceptions, language modules, explicit exclusions, and unknown module codes.",
    )
    add_paragraph(
        doc,
        future_heading,
        "apps/api/src/tests/nusmodsImportService.test.ts verifies module-code, title, units, department, faculty, description, prerequisite text, and prereqTree normalisation, including invalid module-credit fallback and detail-payload merging. These tests isolate deterministic calculations and transformations before they are used by API routes or React pages.",
    )
    add_paragraph(doc, future_heading, "Authorization, Schema, and Service Contract Testing", "h2")
    add_paragraph(
        doc,
        future_heading,
        "apps/api/src/tests/authMiddleware.test.ts verifies unauthenticated rejection, student rejection from administrator routes, and administrator access. apps/api/src/tests/adminCurriculumService.test.ts verifies consecutive academic-year cohorts, normalised and de-duplicated tags, staged mapping deletion, clone self-target rejection, valid drafts, duplicate rule ids, missing rule fields, unknown module references, duplicate mapping changes, and the controlled exception for historical rule modules during clone validation.",
    )
    add_paragraph(doc, future_heading, "Integration Testing", "h2")
    add_paragraph(
        doc,
        future_heading,
        "API integration testing verifies the full route-service-schema path for profile onboarding, multiple plan creation and ordering, owner-scoped updates and deletion, plan import/export, requirement evaluation, settings updates, administrator validation and publication, and curriculum cloning. The checks confirm authentication and administrator authorization, Zod request validation, service composition, persistence semantics, and the response shapes consumed by React Query.",
    )
    add_paragraph(doc, future_heading, "Automated User Interface Testing", "h2")
    add_paragraph(
        doc,
        future_heading,
        "Frontend workflow tests exercise the planner actions proposed in the Milestone 2 README: adding a module, adding a placeholder, moving an item between semesters, renaming and reordering plans, changing the primary plan, importing and exporting with or without grades, deleting a plan, displaying prerequisite warnings, and refreshing degree progress. Milestone 3 workflows also cover viewing a curriculum, editing rule JSON, adding a type-templated rule, staging or deleting a module mapping, validating a draft, publishing a version, creating a curriculum, cloning a cohort, and restricting onboarding or settings to published curricula.",
    )
    add_paragraph(doc, future_heading, "Curriculum Data and Regression Testing", "h2")
    add_paragraph(
        doc,
        future_heading,
        "Fifteen curriculum-data tests cover Computer Science, Business Analytics, Business Artificial Intelligence Systems, Computer Science and Mathematics Double Major, and Business Analytics and Economics Double Degree. They assert the supported cohort, total units, rule types, official core or elective structures, industry and honours pathways, generated classifications, and representative mappings. The evaluator's fixture plans then catch broken allocation logic when a rule type or mapping changes.",
    )
    add_paragraph(doc, future_heading, "GPA and S/U Testing", "h2")
    add_paragraph(
        doc,
        future_heading,
        "GPA tracker testing covers grade-point conversion, pre-S/U and post-S/U GPA, semester GPA, current GPA before the current semester, entire-plan GPA, S/U-only grades, CS grades, zero graded units, placeholders, and missing grades. The S/U eligibility unit tests additionally confirm that the checkbox logic follows module level, NUS-course prerequisites, published Level-2000 exceptions, language-course eligibility, and explicit exclusions.",
    )
    add_paragraph(doc, future_heading, "User Testing", "h2")
    add_paragraph(
        doc,
        future_heading,
        "After scaling to the five currently available programme identities, the team deployed all supported curricula live and invited 20 friends studying one of these programmes to use the application with their own academic plans. The participants tested onboarding, programme and cohort selection, module planning, requirement progress, GPA tracking, S/U decisions, and the general usability of the deployed workflow. Their feedback led to small interface and validation improvements before the Milestone 3 release.",
    )
    add_paragraph(
        doc,
        future_heading,
        "One concrete issue found during this round was that the GPA page originally allowed any module to be marked S/U. Under NUS policy, only eligible modules can be S/Ued. The application now loads module metadata and calls getSuEligibility using the student's cohort. The logic considers course level, whether a Level-2000 course has an NUS-course prerequisite, published exceptions, language-course eligibility, and explicit exclusions. Ineligible modules have a disabled S/U control and an explanatory tooltip. The new policy behaviour is protected by seven dedicated unit tests.",
    )

    replace_paragraph_text(future_heading, "Future Plans")
    database_heading = find_paragraph(doc, "Database Schema Appendix")
    delete_between(future_heading, database_heading)
    add_paragraph(
        doc,
        database_heading,
        "1. Scale from School of Computing to the Entire NUS Population",
        "h2",
    )
    add_paragraph(
        doc,
        database_heading,
        "The next expansion should introduce a faculty-owned curriculum registry. RequirementSet metadata would gain a homeFaculty classification and optional jointFaculty metadata, the administrator catalogue would be filterable by faculty, and onboarding would follow Faculty -> Programme -> Cohort. Shared NUS-wide requirements could still reuse common rule types, while each home faculty would own its source notes, mappings, and regression fixtures. A tangible first phase is to add the metadata and administration filters, onboard at least one complete pilot curriculum outside SoC, validate it against representative plans, and then repeat the same publication workflow faculty by faculty until the full NUS population is covered.",
    )
    add_paragraph(
        doc,
        database_heading,
        "This classification also handles joint programmes cleanly. A double degree remains one curriculum identity, but its primary home faculty and collaborating faculty can be recorded for discovery and maintenance responsibility. The planner and rules engine continue to consume programme/cohort data, while the new faculty layer improves navigation, governance, and scaling.",
    )
    add_paragraph(
        doc,
        database_heading,
        "2. Persisted Curriculum Audit, Review, and Rollback",
        "h2",
    )
    add_paragraph(
        doc,
        database_heading,
        "The current administrator workflow has local drafts, validation, diffs, transactions, and versioned RequirementSet documents. The next tangible step is a persisted audit and approval layer: store draft ownership, source references, change summaries, validation results, reviewer approval, publication timestamp, and the RequirementSet version created by the release. High-impact changes could require a second administrator before publication, and any prior version could be selected as the source of a rollback draft. This would turn technical versioning into an operational curriculum-governance process suitable for multiple faculty maintainers.",
    )
    add_paragraph(
        doc,
        database_heading,
        "3. Constraint-Aware Semester Recommendations",
        "h2",
    )
    add_paragraph(
        doc,
        database_heading,
        "The planner should progress from showing missing requirements to proposing feasible ways to satisfy them. A recommendation service can combine the current RequirementSet, allocation result, prerequisite trees, historical semester offerings, timetable and examination clashes, and a user-selected workload range. It would rank module combinations that close the most requirement gaps without violating prerequisite order or semester constraints, present the reasons for each recommendation, and let the student preview the suggested changes in a new plan before accepting them. A first release can target one empty semester and a single active plan, then expand to multi-semester optimisation.",
    )

    delete_after_to_section_end(database_heading)
    database_content_anchor = doc.add_paragraph(style="normal")
    database_content_anchor.paragraph_format.space_after = Pt(0)
    database_heading._p.addnext(database_content_anchor._p)
    database_heading = database_content_anchor
    add_paragraph(
        doc,
        database_heading,
        "This appendix describes the current Milestone 3 MongoDB models, embedded plan shapes, versioned curriculum records, and non-persisted API payloads. Mongoose defines persistence and indexes in apps/api, while Zod schemas in packages/shared validate API contracts, curriculum rules, administrator drafts, exports, and the discriminated plan-item union. Plan semesters and RequirementSet rules are stored as mixed embedded data in MongoDB because the shared Zod schemas are the canonical shape.",
    )
    add_picture(doc, database_heading, ERD_PATH, width=6.45)
    add_paragraph(
        doc,
        database_heading,
        "Milestone 3 Entity Relationship Diagram (ERD)",
        "caption",
    )
    add_paragraph(doc, database_heading, "MongoDB Collections and API Payloads", "h2")
    add_table(
        doc,
        database_heading,
        ("Model or payload", "Fields and design notes"),
        [
            ("User", "_id, googleId, email, name, avatarUrl?, role, createdAt, updatedAt. role is student or admin; Google identity remains separate from academic state."),
            ("StudentProfile", "_id, userId, programme, cohort, startingSemester, currentSemester, graduationSemester, primaryPlanId?, planOrder[], createdAt, updatedAt. One profile owns the account-wide academic context and plan ordering."),
            ("Plan", "_id, userId, name, programme, cohort, semesters[], createdAt, updatedAt. A user can own many independent scenarios; grades and S/U choices are stored inside each plan's items."),
            ("Module", "_id, acadYear, moduleCode, title, units, department?, faculty?, description?, prerequisite?, prereqTree?, createdAt, updatedAt. This is the shared NUS module catalogue."),
            ("ModuleRequirementTags", "_id, programme, cohort, moduleCode, tags[], createdAt, updatedAt. This maps a shared module to curriculum-specific classifications."),
            ("RequirementSet", "_id, programme, cohort, version, totalUnits, sourceNote, rules[], createdAt, updatedAt. Publication inserts a new version rather than overwriting the previous ruleset."),
            ("PlanExport API payload", "schemaVersion, exportedAt, profile snapshot, and plan snapshot without MongoDB ids. It is encoded as a tcp1_ payload or import link and creates an independent imported Plan."),
            ("AdminCurriculumDraft API payload", "programme, cohort, baseVersion, totalUnits, sourceNote, rules[], tagChanges[]. It is local until validation and is never stored as a MongoDB collection."),
            ("AdminCloneCurriculum API payload", "sourceProgramme, sourceCohort, targetProgramme, targetCohort. Preview checks conflicts and warnings; clone creates target version 1 and copies all module mappings."),
        ],
    )
    add_paragraph(doc, database_heading, "Embedded Plan and Administrator Object Shapes", "h2")
    add_table(
        doc,
        database_heading,
        ("Embedded object", "Shape"),
        [
            ("SemesterPlan", "key, label, items[]. key is IBLOC or Y1S1 through Y5S2 within the profile's starting/graduation range."),
            ("ModulePlanItem", "id?, type: module, moduleCode, units, status, grade?, isSu. status is completed, current, or planned."),
            ("PlaceholderPlanItem", "id?, type: placeholder, requirementId, label, units, grade?, isSu. It reserves requirement units without pretending to be a real module."),
            ("AdminModuleTagChange", "moduleCode, tags[], deleteMapping?. Tags are trimmed, lower-cased, format-validated, and de-duplicated. A deletion must have an empty tag list."),
        ],
    )
    add_paragraph(doc, database_heading, "Indexes and Relations", "h2")
    add_table(
        doc,
        database_heading,
        ("Relation or index", "Purpose"),
        [
            ("User.googleId and User.email are unique", "Prevents duplicate accounts for the same Google identity or email."),
            ("StudentProfile.userId is unique", "Enforces one academic profile per User."),
            ("StudentProfile.primaryPlanId and planOrder[] reference Plan", "Stores the default plan and display order; profile services sanitise both fields to owned plan ids."),
            ("Plan.userId index", "Supports efficient owner-scoped plan listing and mutation."),
            ("Module unique acadYear + moduleCode", "Prevents duplicate catalogue entries for one module in the same academic year."),
            ("ModuleRequirementTags unique programme + cohort + moduleCode", "Allows exactly one mapping document per module for each curriculum identity."),
            ("RequirementSet unique programme + cohort + version", "Preserves append-only curriculum versions while rejecting duplicate version numbers."),
            ("Plan.programme + cohort resolves RequirementSet and mappings", "Evaluation loads the latest matching RequirementSet and programme/cohort-specific tag lookup."),
        ],
    )
    add_paragraph(doc, database_heading, "Requirement Rule Shape", "h2")
    add_table(
        doc,
        database_heading,
        ("Rule object", "Supported fields"),
        [
            ("RequirementSet", "programme, cohort, version, totalUnits, sourceNote, rules[]."),
            ("Common RequirementRule fields", "id, label, type, requiredUnits?, requiredModules?, moduleOptions?, acceptedTags?, acceptedPlaceholders?, tagCaps?, advisory?."),
            ("Structured ID/CD and breadth/depth fields", "idTags?, cdTags?, requiredIdMinCourses?, allowedCdMaxCourses?, focusAreas?, focus-area counts, Level-4000 units, industry min/max, allowed prefixes, CP cap, industryTags?, dissertationTags?."),
            ("Programme elective and industry fields", "requiredMinCourses?, requiredLevel4000MinCourses?, requiredPrefixMinCourses?, requiredPrefixes?, internshipFoundationTags?, secondInternshipTags?, supplementaryTags?, requiredFoundationUnits?, requiredCompanionUnits?, tagUnitOverrides?."),
            ("Double-degree honours fields", "integratedThesisTags?, economicsElectiveTags?, economicsLevel4000Tags?, integratedThesisUnits?, integratedEconomicsUnits?, integratedEconomicsLevel4000Units?, internshipEconomicsUnits?, internshipEconomicsLevel4000Units?, internshipPathwayRequiredUnits?."),
            ("Rule types", "module-list, module-choice, units-from-tags, capped-units-from-tags, placeholder-units, combined-units, structured-idcd, structured-breadth-depth, structured-programme-electives, structured-industry-experience, structured-ddp-honours-pathway, residual-units."),
        ],
    )
    add_paragraph(doc, database_heading, "Administrator Draft and Publication Semantics", "h2")
    add_paragraph(
        doc,
        database_heading,
        "AdminCurriculumDraft is a non-persisted API contract with baseVersion. Validation overlays tagChanges on the current mapping lookup, verifies rule semantics and module references, compares baseVersion with the latest RequirementSet, and evaluates up to five recent plans. Publication repeats the version check inside a MongoDB transaction, inserts the next RequirementSet version, and applies mapping upserts or deletions atomically.",
    )
    add_paragraph(
        doc,
        database_heading,
        "Clone preview is also non-persisted. It checks that the target RequirementSet and mapping namespace are empty, retains historical rule references with warnings, and reports source statistics. A successful clone creates target version 1 and copies every source ModuleRequirementTags record in the same transaction.",
    )
    add_paragraph(doc, database_heading, "Why These Relations Matter", "h2")
    add_paragraph(
        doc,
        database_heading,
        "The schema keeps identity, academic profile, planning scenarios, shared module metadata, curriculum-specific module interpretation, and versioned curriculum rules separate. Each category changes for a different reason: users update timelines, plans branch into alternatives, NUSMods updates catalogue metadata, and faculties publish new curriculum versions.",
    )
    add_paragraph(
        doc,
        database_heading,
        "Plan embeds SemesterPlan and PlanItem because planner and GPA screens usually read and save the complete scenario. Module and RequirementSet are shared reference data. ModuleRequirementTags bridges a module's stable identity to a curriculum's meaning. PlanExport, AdminCurriculumDraft, and AdminCloneCurriculum remain API payloads rather than collections so sharing creates independent copies and administration publishes only validated versioned records.",
    )
    add_paragraph(
        doc,
        database_heading,
        "Versioned RequirementSet documents plus transactional mapping publication provide curriculum history without duplicating the shared module catalogue. The programme/cohort identity also lets onboarding, settings, plan evaluation, administrator selection, and test fixtures resolve the same curriculum consistently.",
    )

    doc.core_properties.title = "MS3 README"
    doc.core_properties.subject = "The CS Plan - NUS Orbital 2026 Milestone 3"
    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
